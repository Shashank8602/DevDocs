"""Chunkers: every line range must match the source lines, pages/sections must be recorded."""
import io
import json

import pytest
from docx import Document

from app.ingestion.chunking import MAX_CHARS, chunk_document
from app.ingestion.parsers import ParseError, parse

PY = '''import os

TOKEN = "x"


# Verify a login.
@decorator
def login(user):
    return user


class Auth:
    def check(self):
        return True

    def refresh(self):
        return False


async def logout():
    pass
'''


def chunks_for(path: str, text: str | bytes, file_type="source", language=None):
    data = text if isinstance(text, bytes) else text.encode()
    return chunk_document(parse(data, path, file_type, language))


def assert_lines_match(source: str, chunks) -> None:
    lines = source.split("\n")
    for c in chunks:
        assert c.line_start is not None and c.line_end is not None
        got = [x.strip() for x in lines[c.line_start - 1 : c.line_end]]
        assert got == [x.strip() for x in c.content.split("\n")], f"chunk {c.ordinal} does not match its lines"


def make_pdf(pages: list[str]) -> bytes:
    """A minimal valid PDF with one text page per entry."""
    objs: list[bytes] = [b"<< /Type /Catalog /Pages 2 0 R >>"]
    kids = " ".join(f"{3 + 2 * i} 0 R" for i in range(len(pages)))
    objs.append(f"<< /Type /Pages /Kids [{kids}] /Count {len(pages)} >>".encode())
    font_id = 3 + 2 * len(pages)
    for i, text in enumerate(pages):
        page_id = 3 + 2 * i
        stream = "BT /F1 10 Tf 40 750 Td 12 TL " + " ".join(
            f"({line.replace('(', '[').replace(')', ']')}) '" for line in text.split("\n")
        ) + " ET"
        objs.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {page_id + 1} 0 R "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> >>".encode()
        )
        objs.append(f"<< /Length {len(stream)} >>\nstream\n{stream}\nendstream".encode())
    objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for n, body in enumerate(objs, start=1):
        offsets.append(out.tell())
        out.write(f"{n} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objs) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objs) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def test_python_splits_at_function_and_class_boundaries():
    chunks = chunks_for("auth.py", PY, language="python")
    assert_lines_match(PY.rstrip("\n"), chunks)
    sections = [c.section for c in chunks]
    assert sections == [None, "login", "Auth", "logout"]
    login = chunks[1]
    assert login.content.startswith("# Verify a login.\n@decorator\ndef login")  # comment + decorator included
    assert (login.line_start, login.line_end) == (6, 9)


def test_long_python_class_is_split_by_method():
    body = "\n".join(f"    def m{i}(self):\n        return {i}\n" for i in range(40))
    src = f"class Big:\n{body}"
    chunks = chunks_for("big.py", src, language="python")
    assert_lines_match(src.rstrip("\n"), chunks)
    assert {c.section for c in chunks} >= {"Big.m0", "Big.m39"}


def test_python_syntax_error_falls_back_to_generic_splitting():
    chunks = chunks_for("bad.py", "def broken(:\n    pass\n", language="python")
    assert len(chunks) == 1 and chunks[0].line_start == 1


def test_oversized_python_function_is_split_with_exact_lines():
    src = "def big():\n" + "".join(f"    value_{i} = compute({i})  # step {i}\n" for i in range(400))
    chunks = chunks_for("big.py", src, language="python")
    assert len(chunks) > 1
    assert all(len(c.content) <= MAX_CHARS for c in chunks)
    assert_lines_match(src.rstrip("\n"), chunks)
    assert chunks[0].line_start == 1 and chunks[-1].line_end == 401


def test_other_languages_use_language_aware_splitting_with_lines():
    src = "".join(f"function f{i}() {{\n  return {i};\n}}\n\n" for i in range(300))
    chunks = chunks_for("app.js", src, language="javascript")
    assert len(chunks) > 1
    assert_lines_match(src.rstrip("\n"), chunks)


def test_markdown_splits_by_heading_with_path_and_ignores_headings_in_code_fences():
    md = "intro\n\n# Setup\n\ntext\n\n## Install\n\n```bash\n# not a heading\npip install x\n```\n\n# Usage\n\nrun it\n"
    chunks = chunks_for("README.md", md, "doc", "markdown")
    assert [c.section for c in chunks] == [None, "Setup", "Setup > Install", "Usage"]
    assert_lines_match(md.rstrip("\n"), chunks)
    assert "# not a heading" in chunks[2].content


def test_pdf_chunks_by_page_and_oversized_page_keeps_its_page_number():
    long_page = "\n".join(f"Line {i} describes how the payment service retries failed charges." for i in range(120))
    data = make_pdf(["First page about auth.", "Second page.", long_page])
    chunks = chunks_for("docs/a.pdf", data, "doc")
    assert [c.page_number for c in chunks][:2] == [1, 2]
    big = [c for c in chunks if c.page_number == 3]
    assert len(big) >= 2
    assert all(len(c.content) <= MAX_CHARS and c.line_start is None for c in big)
    assert [c.ordinal for c in chunks] == list(range(len(chunks)))


def test_corrupt_pdf_and_textless_pdf_raise_parse_error():
    with pytest.raises(ParseError):
        parse(b"this is not really a pdf", "x.pdf", "doc", None)
    with pytest.raises(ParseError):
        parse(make_pdf([""]), "blank.pdf", "doc", None)


def test_docx_splits_by_heading_and_includes_tables():
    doc = Document()
    doc.add_heading("Requirements", 1)
    doc.add_paragraph("Users can log in.")
    doc.add_heading("Payments", 2)
    doc.add_paragraph("Cards are charged once.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Plan", "Price"
    table.cell(1, 0).text, table.cell(1, 1).text = "Pro", "10"
    buf = io.BytesIO()
    doc.save(buf)
    chunks = chunks_for("req.docx", buf.getvalue(), "doc")
    assert [c.section for c in chunks] == ["Requirements", "Requirements > Payments"]
    assert "Plan | Price" in chunks[1].content and "Pro | 10" in chunks[1].content
    assert chunks[0].line_start is None


def test_json_is_split_at_top_level_keys_and_small_keys_are_packed():
    src = json.dumps({"name": "x", "scripts": {"a": "1"}, "deps": {"b": "2"}}, indent=2)
    chunks = chunks_for("package.json", src, "other", "json")
    assert len(chunks) == 1  # tiny keys share a chunk
    assert (chunks[0].line_start, chunks[0].line_end) == (1, len(src.split("\n")))
    big = json.dumps({"a": {f"k{i}": "v" * 40 for i in range(80)}, "b": {f"k{i}": "w" * 40 for i in range(80)}}, indent=2)
    chunks = chunks_for("big.json", big, "other", "json")
    assert len(chunks) > 1
    assert_lines_match(big, chunks)


def test_invalid_json_still_becomes_text_chunks():
    assert len(chunks_for("bad.json", '{"a": ', "other", "json")) == 1


def test_yaml_sections_use_top_level_keys():
    src = "server:\n  port: 8000\n\ndatabase:\n  url: x\n"
    chunks = chunks_for("config.yaml", src, "other", "yaml")
    assert len(chunks) == 1 and chunks[0].section == "server – database"
    assert (chunks[0].line_start, chunks[0].line_end) == (1, 5)
    assert_lines_match(src.rstrip("\n"), chunks)


def test_csv_groups_rows_repeats_header_and_tracks_lines():
    rows = [f"{i},product {i},{i * 2}.50" for i in range(1, 400)]
    src = "id,name,price\n" + "\n".join(rows) + "\n"
    chunks = chunks_for("products.csv", src, "other")
    assert len(chunks) > 1
    assert all(c.content.startswith("id,name,price\n") for c in chunks)
    assert chunks[0].line_start == 2 and chunks[-1].line_end == 400
    lines = src.split("\n")
    for c in chunks:
        assert c.content == "id,name,price\n" + "\n".join(lines[c.line_start - 1 : c.line_end])
    assert all(a.line_end + 1 == b.line_start for a, b in zip(chunks, chunks[1:]))


def test_csv_with_multiline_quoted_field_keeps_record_together():
    src = 'id,note\n1,"line one\nline two"\n2,plain\n'
    chunks = chunks_for("n.csv", src, "other")
    assert chunks[0].line_start == 2 and chunks[0].line_end == 4


def test_txt_and_html_keep_line_ranges():
    txt = "\n\n".join(f"Paragraph {i}. " + "words " * 60 for i in range(40))
    chunks = chunks_for("notes.txt", txt, "doc")
    assert len(chunks) > 1
    assert_lines_match(txt, chunks)
    html = "<html>\n<body>\n" + "".join(f"<div>\n  <p>item {i}</p>\n</div>\n" for i in range(500)) + "</body>\n</html>"
    chunks = chunks_for("index.html", html, "source", "html")
    assert len(chunks) > 1
    assert_lines_match(html, chunks)


def test_empty_file_has_no_chunks_and_crlf_is_normalised():
    assert chunks_for("empty.txt", "  \n", "doc") == []
    chunks = chunks_for("w.py", "def a():\r\n    pass\r\n", language="python")
    assert chunks[0].content == "def a():\n    pass" and chunks[0].line_end == 2


def test_binary_content_is_rejected():
    with pytest.raises(ParseError):
        parse(b"abc\x00\x01\x02", "x.txt", "doc", None)
